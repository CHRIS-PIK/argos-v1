from __future__ import annotations

import csv
import io
import json
import logging
import os
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator, Literal

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, ConfigDict, Field, field_validator

from app.ingest.keys import CollectorIdentity
from app.query import db
from app.query.auth import require_reader
from app.query.filters import DIM_KEY, MAX_DIMENSIONS, Filters, common_filters
from app.query.sql import (
    count_query,
    entities_query,
    entity_cursor,
    event_cursor,
    events_query,
    resolve_window,
)
from app.queue import get_redis

log = logging.getLogger(__name__)
router = APIRouter()

BATCH = 5000
XLSX_MAX_ROWS = 1_048_575
XLSX_MAX_CELL = 32_767
JOB_PREFIX = "argos:exports:"

CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "jsonl": "application/x-ndjson",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "parquet": "application/vnd.apache.parquet",
}
TIME_COLUMNS = {"bucket_at", "collected_at", "ingested_at"}


def max_rows() -> int:
    return int(os.getenv("EXPORT_MAX_ROWS", "1000000"))


def direct_limit() -> int:
    return int(os.getenv("EXPORT_DIRECT_LIMIT", "50000"))


def export_dir() -> Path:
    path = Path(os.getenv("EXPORT_DIR", "/tmp/argos-exports"))
    path.mkdir(parents=True, exist_ok=True)
    return path


def job_ttl() -> int:
    return int(os.getenv("EXPORT_TTL_SECONDS", str(24 * 3600)))


class ExportRequest(BaseModel):
    model_config = ConfigDict(populate_by_name=True)

    dataset: Literal["entities", "events"]
    format: Literal["csv", "jsonl", "xlsx", "parquet"] = "csv"
    event_type: str | None = Field(default=None, max_length=64)
    source: str | None = Field(default=None, max_length=64)
    status: str | None = Field(default=None, max_length=128)
    q: str | None = Field(default=None, max_length=128)
    entity_id: str | None = Field(default=None, max_length=512)
    dimensions: dict[str, str] = Field(default_factory=dict)
    time_from: datetime | None = Field(default=None, alias="from")
    time_to: datetime | None = Field(default=None, alias="to")
    include_payload: bool = False
    dimension_columns: list[str] = Field(default_factory=list, max_length=20)

    @field_validator("dimensions")
    @classmethod
    def _dimensions(cls, value: dict[str, str]) -> dict[str, str]:
        if len(value) > MAX_DIMENSIONS or any(not DIM_KEY.match(k) for k in value):
            raise ValueError("invalid dimension filters")
        return value

    @field_validator("dimension_columns")
    @classmethod
    def _dimension_columns(cls, value: list[str]) -> list[str]:
        if any(not DIM_KEY.match(k) for k in value):
            raise ValueError("invalid dimension column name")
        return value

    def filters(self) -> Filters:
        return Filters(
            event_type=self.event_type, source=self.source, status=self.status, q=self.q,
            entity_id=self.entity_id if self.dataset == "events" else None,
            dimensions=self.dimensions,
        )


def columns_for(req: ExportRequest) -> list[str]:
    columns = ["tenant_id", "event_type", "source", "schema_version", "entity_id"]
    if req.dataset == "events":
        columns.append("bucket_at")
    columns += ["name", "status", "collected_at", "ingested_at", "batch_id"]
    columns += [f"dim.{k}" for k in req.dimension_columns] if req.dimension_columns else ["dimensions"]
    if req.include_payload:
        columns.append("payload")
    return columns


def _aware(value: datetime) -> datetime:
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def _iso(value: datetime) -> str:
    return _aware(value).isoformat().replace("+00:00", "Z")


def flatten(row: dict[str, Any], columns: list[str], native_time: bool = False) -> dict[str, Any]:
    out: dict[str, Any] = {}
    dimensions = row.get("dimensions") or {}
    for column in columns:
        if column.startswith("dim."):
            out[column] = dimensions.get(column[4:], "")
        elif column == "dimensions":
            out[column] = json.dumps(dimensions, ensure_ascii=False, separators=(",", ":"))
        elif column in TIME_COLUMNS:
            value = row.get(column)
            out[column] = None if value is None else (_aware(value) if native_time else _iso(value))
        else:
            out[column] = row.get(column, "")
    return out


def iter_rows(req: ExportRequest, identity: CollectorIdentity, limit_total: int) -> Iterator[dict[str, Any]]:
    """Percorre o resultado em lotes por keyset, sem OFFSET."""
    filters = req.filters()
    window = resolve_window(req.time_from, req.time_to) if req.dataset == "events" else None
    cursor: list[Any] | None = None
    emitted = 0
    while emitted < limit_total:
        size = min(BATCH, limit_total - emitted)
        if req.dataset == "entities":
            sql, params = entities_query(filters, identity, size, cursor, req.include_payload)
            cursor_of = entity_cursor
        else:
            sql, params = events_query(filters, identity, *window, size, cursor, req.include_payload)
            cursor_of = event_cursor
        batch = db.rows(sql, params)
        yield from batch
        emitted += len(batch)
        if len(batch) < size:
            return
        cursor = cursor_of(batch[-1])


def count_rows(req: ExportRequest, identity: CollectorIdentity) -> int:
    window = resolve_window(req.time_from, req.time_to) if req.dataset == "events" else None
    sql, params = count_query(req.dataset, req.filters(), identity, window)
    return int(db.rows(sql, params)[0]["total"])


# ---------------------------------------------------------------- escritores


def text_chunks(fmt: str, columns: list[str], rows: Iterator[dict[str, Any]]) -> Iterator[str]:
    """CSV e JSON Lines em pedaços de texto, serve tanto para streaming quanto para arquivo."""
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(columns)
        yield buffer.getvalue()
        buffer.seek(0), buffer.truncate()
        count = 0
        for row in rows:
            flat = flatten(row, columns)
            writer.writerow([flat[c] for c in columns])
            count += 1
            if count % 1000 == 0:
                yield buffer.getvalue()
                buffer.seek(0), buffer.truncate()
        if buffer.tell():
            yield buffer.getvalue()
    else:
        lines = []
        for row in rows:
            lines.append(json.dumps(flatten(row, columns), ensure_ascii=False))
            if len(lines) >= 1000:
                yield "\n".join(lines) + "\n"
                lines = []
        if lines:
            yield "\n".join(lines) + "\n"


def write_file(path: Path, fmt: str, columns: list[str], rows: Iterator[dict[str, Any]]) -> int:
    """Grava o arquivo e devolve o número de linhas escritas."""
    written = 0
    if fmt in ("csv", "jsonl"):
        counted = _counting(rows)
        with path.open("w", encoding="utf-8-sig" if fmt == "csv" else "utf-8", newline="") as fp:
            for chunk in text_chunks(fmt, columns, counted):
                fp.write(chunk)
        return counted.count

    if fmt == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("argos")
        sheet.append(columns)
        for row in rows:
            if written >= XLSX_MAX_ROWS:
                raise ValueError("result exceeds the Excel row limit; use csv or parquet")
            flat = flatten(row, columns)
            sheet.append([v[:XLSX_MAX_CELL] if isinstance(v, str) else v for v in flat.values()])
            written += 1
        workbook.save(path)
        return written

    import pyarrow as pa
    import pyarrow.parquet as pq

    schema = pa.schema([
        (c, pa.timestamp("ms", tz="UTC") if c in TIME_COLUMNS else pa.string()) for c in columns
    ])
    with pq.ParquetWriter(path, schema) as writer:
        batch: list[dict[str, Any]] = []
        for row in rows:
            batch.append(flatten(row, columns, native_time=True))
            written += 1
            if len(batch) >= BATCH:
                writer.write_table(pa.Table.from_pylist(batch, schema=schema))
                batch = []
        if batch:
            writer.write_table(pa.Table.from_pylist(batch, schema=schema))
    return written


class _counting:
    def __init__(self, rows: Iterator[dict[str, Any]]) -> None:
        self._rows = rows
        self.count = 0

    def __iter__(self):
        for row in self._rows:
            self.count += 1
            yield row


# ---------------------------------------------------------------- jobs

executor = ThreadPoolExecutor(max_workers=int(os.getenv("EXPORT_WORKERS", "2")))


def _job_key(job_id: str) -> str:
    return JOB_PREFIX + job_id


def _sweep_old_files() -> None:
    cutoff = time.time() - job_ttl()
    for path in export_dir().glob("*"):
        try:
            if path.is_file() and path.stat().st_mtime < cutoff:
                path.unlink()
        except OSError:
            pass


def run_job(job_id: str, req: ExportRequest, identity: CollectorIdentity) -> None:
    client = get_redis()
    key = _job_key(job_id)
    path = export_dir() / f"{job_id}.{req.format}"
    try:
        client.hset(key, mapping={"status": "running"})
        total = count_rows(req, identity)
        limit = max_rows()
        rows = write_file(path, req.format, columns_for(req), iter_rows(req, identity, limit))
        client.hset(key, mapping={
            "status": "done", "rows": rows, "total": total, "truncated": int(total > limit),
            "file": path.name, "finished_at": datetime.now(timezone.utc).isoformat(),
        })
    except Exception as exc:
        log.exception("export failed job=%s", job_id)
        path.unlink(missing_ok=True)
        client.hset(key, mapping={
            "status": "failed", "error": str(exc)[:500],
            "finished_at": datetime.now(timezone.utc).isoformat(),
        })
    client.expire(key, job_ttl())


def create_job(req: ExportRequest, identity: CollectorIdentity) -> str:
    _sweep_old_files()
    job_id = uuid.uuid4().hex
    client = get_redis()
    client.hset(_job_key(job_id), mapping={
        "status": "queued", "tenant_id": identity.tenant_id, "dataset": req.dataset,
        "format": req.format, "created_at": datetime.now(timezone.utc).isoformat(),
    })
    client.expire(_job_key(job_id), job_ttl())
    executor.submit(run_job, job_id, req, identity)
    return job_id


def load_job(job_id: str, identity: CollectorIdentity) -> dict[str, str]:
    record = get_redis().hgetall(_job_key(job_id))
    if not record or record.get("tenant_id") != identity.tenant_id:
        raise HTTPException(status_code=404, detail="export not found")
    return record


def public_job(job_id: str, record: dict[str, str]) -> dict[str, Any]:
    out: dict[str, Any] = {"id": job_id, **{k: v for k, v in record.items() if k not in ("tenant_id", "file")}}
    for numeric in ("rows", "total"):
        if numeric in out:
            out[numeric] = int(out[numeric])
    if "truncated" in out:
        out["truncated"] = out["truncated"] == "1"
    if record.get("status") == "done":
        out["download_url"] = f"/v1/exports/{job_id}/download"
    return out


# ---------------------------------------------------------------- endpoints


@router.get("/v1/export")
def export_direct(
    dataset: Literal["entities", "events"],
    format: Literal["csv", "jsonl"] = "csv",
    filters: Filters = Depends(common_filters),
    identity: CollectorIdentity = Depends(require_reader),
    entity_id: str | None = Query(default=None, max_length=512),
    time_from: datetime | None = Query(default=None, alias="from"),
    time_to: datetime | None = Query(default=None, alias="to"),
    include_payload: bool = False,
    dimension_columns: str = Query(default="", description="Chaves separadas por vírgula"),
):
    """Download direto em streaming, para resultados pequenos. Acima do limite, use POST /v1/exports."""
    req = ExportRequest(
        dataset=dataset, format=format, event_type=filters.event_type, source=filters.source,
        status=filters.status, q=filters.q, entity_id=entity_id, dimensions=filters.dimensions,
        time_from=time_from, time_to=time_to, include_payload=include_payload,
        dimension_columns=[k for k in dimension_columns.split(",") if k],
    )
    total = count_rows(req, identity)
    if total > direct_limit():
        raise HTTPException(
            status_code=413,
            detail=f"{total} rows exceed the direct download limit ({direct_limit()}); "
                   "create an export with POST /v1/exports",
        )
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    return StreamingResponse(
        text_chunks(format, columns_for(req), iter_rows(req, identity, direct_limit())),
        media_type=CONTENT_TYPES[format],
        headers={"Content-Disposition": f'attachment; filename="argos-{dataset}-{stamp}.{format}"'},
    )


@router.post("/v1/exports", status_code=202)
def create_export(req: ExportRequest, identity: CollectorIdentity = Depends(require_reader)):
    if req.dataset == "events":
        resolve_window(req.time_from, req.time_to)
    job_id = create_job(req, identity)
    return public_job(job_id, load_job(job_id, identity))


@router.get("/v1/exports/{job_id}")
def get_export(job_id: str, identity: CollectorIdentity = Depends(require_reader)):
    return public_job(job_id, load_job(job_id, identity))


@router.get("/v1/exports/{job_id}/download")
def download_export(job_id: str, identity: CollectorIdentity = Depends(require_reader)):
    record = load_job(job_id, identity)
    if record.get("status") != "done":
        raise HTTPException(status_code=409, detail=f"export is {record.get('status')}")
    path = export_dir() / record["file"]
    if not path.is_file():
        raise HTTPException(status_code=410, detail="export file expired")
    fmt = record["format"]
    return FileResponse(
        path, media_type=CONTENT_TYPES[fmt],
        filename=f"argos-{record['dataset']}-{job_id[:8]}.{fmt}",
    )
